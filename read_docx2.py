import sys
import docx

def read_docx(path, out_path):
    try:
        doc = docx.Document(path)
        full_text = []
        for para in doc.paragraphs:
            if para.text.strip():
                full_text.append(para.text)
        
        for table in doc.tables:
            for row in table.rows:
                row_data = []
                for cell in row.cells:
                    if cell.text.strip():
                        # remove newlines within cell
                        row_data.append(cell.text.replace('\n', ' ').strip())
                if row_data:
                    full_text.append(" | ".join(row_data))
                    
        with open(out_path, 'w', encoding='utf-8') as f:
            f.write('\n'.join(full_text))
        print("Success")
    except Exception as e:
        print(str(e))

read_docx(r'D:\Binh_KL\Code\Thesis_Proposal_NguyenThanhBinh.docx', r'D:\Binh_KL\Code\extracted_text.txt')
